"""Document loading: recursively scan directories for .md/.txt (and .pdf when
the optional pypdf package is installed), peel off Obsidian-style YAML
frontmatter (tags/aliases subset), collect [[wikilinks]], and expand chat-log
exports (ChatGPT / Claude `conversations.json`) into per-conversation docs.
.html/.htm files are read with the stdlib html.parser — no extra dependency."""
from __future__ import annotations

import hashlib
import re
from html.parser import HTMLParser
from pathlib import Path

from loci.chatlog import CHATLOG_NAMES, parse_chatlog

try:
    from pypdf import PdfReader
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

try:
    import pymupdf4llm
    HAS_PDF_TABLES = True
except ImportError:
    HAS_PDF_TABLES = False

try:
    import docx as _docx  # python-docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

SUFFIXES = {".md", ".txt"}
HTML_SUFFIXES = {".html", ".htm"}
ORG_SUFFIXES = {".org", ".org_archive"}
_WIKILINK = re.compile(r"(?<!!)\[\[([^\]|#]+)(?:[#|][^\]]*)?\]\]")
_ORG_LINK = re.compile(r"\[\[([^\]]+)\](?:\[([^\]]*)\])?\]")


def _read_org(text: str) -> str:
    """Convert org-mode markup to markdown-ish text.

    `*`-headings become `#`-headings, #+FILETAGS becomes YAML frontmatter tags
    (so tags_of() picks them up), #+TITLE becomes an h1, and org links
    [[url][desc]] / [[url]] become markdown links. Everything else passes
    through untouched — org bodies are already plain prose."""
    frontmatter: list[str] = []
    title = ""
    lines: list[str] = []
    # org semantics: #+TITLE is the document title and `*` sections nest under
    # it — so when a title exists, headings shift down one level to keep the
    # hierarchy faithful (TITLE→h1, `*`→h2, `**`→h3).
    has_title = bool(re.search(r"^#\+TITLE:\s*\S", text, re.MULTILINE))
    for line in text.splitlines():
        m = re.match(r"^#\+TITLE:\s*(.*)$", line)
        if m:
            title = m.group(1).strip()
            continue
        m = re.match(r"^#\+FILETAGS:\s*:?(.*?)\s*:?$", line)
        if m:
            tags = [t for t in m.group(1).split(":") if t]
            if tags:
                frontmatter.append("tags: [" + ", ".join(tags) + "]")
            continue
        if re.match(r"^#\+[A-Z_]+:", line):     # other #+KEYWORDS: drop
            continue
        m = re.match(r"^(\*+)\s+(.*)$", line)
        if m:
            level = min(len(m.group(1)) + (1 if has_title else 0), 6)
            lines.append("#" * level + " " + m.group(2).strip())
            continue
        m = re.match(r"^(\s*)\+\s+(.*)$", line)  # org's alternate "+ " bullet
        if m:
            lines.append(m.group(1) + "- " + m.group(2))
            continue
        lines.append(_ORG_LINK.sub(
            lambda lm: f"[{lm.group(2) or lm.group(1)}]({lm.group(1)})"
            if "://" in lm.group(1) or lm.group(1).startswith("mailto:")
            else (lm.group(2) or lm.group(1)),  # internal links: keep the text
            line))
    body = "\n".join(lines)
    if title:
        body = f"# {title}\n\n{body}"
    if frontmatter:
        body = "---\n" + "\n".join(frontmatter) + "\n---\n\n" + body
    return body


def extract_wikilinks(body: str) -> str:
    """Collect [[wikilink]] targets as a comma-joined string (deduplicated).
    Obsidian embeds (![[...]]) are assets, not links — excluded."""
    seen: list[str] = []
    for m in _WIKILINK.finditer(body):
        target = m.group(1).strip()
        if target and target.lower() not in (t.lower() for t in seen):
            seen.append(target)
    return ",".join(seen)


def file_hash(content: str) -> str:
    return hashlib.sha1(content.encode("utf-8")).hexdigest()


def parse_frontmatter(content: str) -> tuple[dict, str]:
    """Parse a minimal YAML subset from leading --- fences: scalars and lists.
    Returns ({meta}, body); malformed fences are left in the body untouched."""
    if not content.startswith("---"):
        return {}, content
    lines = content.splitlines()
    if lines[0].strip() != "---":
        return {}, content
    end = None
    for i in range(1, len(lines)):
        if lines[i].strip() == "---":
            end = i
            break
    if end is None:
        return {}, content
    meta: dict = {}
    current_list: str | None = None
    for line in lines[1:end]:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("- ") and current_list:
            meta[current_list].append(stripped[2:].strip().strip('"\''))
            continue
        if ":" not in stripped:
            continue
        key, _, value = stripped.partition(":")
        key, value = key.strip(), value.strip()
        current_list = None
        if not value:
            meta[key] = []
            current_list = key
        elif value.startswith("[") and value.endswith("]"):
            inner = value[1:-1].strip()
            meta[key] = [v.strip().strip("\"'") for v in inner.split(",")] if inner else []
        else:
            meta[key] = value.strip("\"'")
    return meta, "\n".join(lines[end + 1:])


def tags_of(meta: dict) -> str:
    """Normalize the frontmatter `tags` key (list or string) to a comma-joined string."""
    raw = meta.get("tags", [])
    if isinstance(raw, str):
        items = [t.strip() for t in raw.split(",")]
    else:
        items = [str(t).strip() for t in raw]
    return ",".join(t for t in items if t)


class _HTMLToText(HTMLParser):
    """Extract readable blocks from HTML into markdown-ish text.

    Headings become `#`-prefixed lines, <li> becomes `- `, paragraph-level
    tags flush the inline buffer into blank-line-separated blocks. Script/
    style content is skipped entirely."""
    BLOCKS = {"p", "div", "section", "article", "aside", "blockquote", "pre",
              "table", "tr", "header", "footer", "nav", "figure", "figcaption",
              "main", "ul", "ol", "dl", "dd", "dt", "hr"}
    HEADINGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
    SKIP = {"script", "style", "noscript", "template", "svg"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._buf: list[str] = []
        self._pending = ""      # prefix queued by the opening tag, spent on flush
        self._skip_depth = 0
        self._title: list[str] = []
        self._in_title = False

    def _flush(self) -> None:
        text = re.sub(r"\s+", " ", "".join(self._buf)).strip()
        prefix, self._pending = self._pending, ""
        self._buf.clear()
        if text:
            self.parts.append(prefix + text)

    def handle_starttag(self, tag, attrs):
        if self._skip_depth:
            if tag in self.SKIP:
                self._skip_depth += 1
            return
        if tag in self.SKIP:
            self._skip_depth += 1
        elif tag == "title":
            self._in_title = True
        elif tag in self.HEADINGS:
            self._flush()
            self._pending = "#" * int(tag[1]) + " "
        elif tag == "li":
            self._flush()
            self._pending = "- "
        elif tag == "br":
            self._buf.append(" ")
        elif tag in self.BLOCKS:
            self._flush()

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self._skip_depth = max(0, self._skip_depth - 1)
            return
        if self._skip_depth:
            return
        if tag == "title":
            self._in_title = False
        elif tag in self.HEADINGS or tag in self.BLOCKS or tag == "li":
            self._flush()

    def handle_data(self, data):
        if self._skip_depth:
            return
        if self._in_title:
            self._title.append(data)
        else:
            self._buf.append(data)

    def to_text(self) -> str:
        self._flush()
        title = re.sub(r"\s+", " ", "".join(self._title)).strip()
        # <title> is often site chrome; only surface it when the body has no
        # headings of its own to anchor the chunk hierarchy.
        if title and not any(p.startswith("#") for p in self.parts):
            parts = [f"# {title}"] + self.parts
        else:
            parts = self.parts
        return "\n\n".join(parts)


def _read_html(p: Path) -> str | None:
    raw = p.read_bytes()
    encoding = "utf-8"
    m = re.search(rb'charset=["\']?([\w-]+)', raw[:2048], re.IGNORECASE)
    if m:
        encoding = m.group(1).decode("ascii", errors="ignore")
    try:
        text = raw.decode(encoding)
    except (UnicodeDecodeError, LookupError):
        try:  # common for Chinese exports without a meta charset
            text = raw.decode("gb18030")
        except UnicodeDecodeError:
            print(f"[warn] html undecodable, skipping: {p.name}")
            return None
    parser = _HTMLToText()
    try:
        parser.feed(text)
    except Exception as e:
        print(f"[warn] html unreadable, skipping: {p.name} ({e})")
        return None
    return parser.to_text()


def _read_text(p: Path) -> str | None:
    """Read a document as text/markdown; returns None when it can't be handled."""
    suffix = p.suffix.lower()
    if suffix in SUFFIXES:
        try:
            return p.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            print(f"[warn] not UTF-8, skipping: {p.name}")
            return None
    if suffix in ORG_SUFFIXES:
        try:
            return _read_org(p.read_text(encoding="utf-8"))
        except UnicodeDecodeError:
            print(f"[warn] not UTF-8, skipping: {p.name}")
            return None
    if suffix in HTML_SUFFIXES:
        return _read_html(p)
    if suffix == ".pdf":
        if HAS_PDF_TABLES:
            try:
                # markdown output: headings, lists, and tables as pipe rows
                return pymupdf4llm.to_markdown(str(p))
            except Exception as e:
                print(f"[warn] pdf table extraction failed, falling back: {p.name} ({e})")
        if HAS_PDF:
            try:
                reader = PdfReader(str(p))
                return "\n\n".join((page.extract_text() or "") for page in reader.pages)
            except Exception as e:
                print(f"[warn] pdf unreadable, skipping: {p.name} ({e})")
                return None
        return None  # no pdf extra installed; `doctor` mentions it
    if suffix == ".docx":
        if not HAS_DOCX:
            return None  # silently skip; `doctor` mentions the optional extra
        try:
            document = _docx.Document(str(p))
            parts = [para.text for para in document.paragraphs if para.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            return "\n\n".join(parts)
        except Exception as e:
            print(f"[warn] docx unreadable, skipping: {p.name} ({e})")
            return None
    return None


def scan_sources(sources: list[dict]) -> list[dict]:
    """Return [{path, content, hash, tags, links, mtime, chunk_size?, chunk_overlap?}].
    Each [[sources]] entry may override `chunk_size` / `chunk_overlap`; absent
    keys mean "use the global [chunk] defaults"."""
    docs, seen = [], set()
    for src in sources:
        root = Path(src["path"]).expanduser()
        if not root.exists():
            print(f"[warn] directory not found, skipping: {root}")
            continue
        src_chunk_size = src.get("chunk_size")
        src_chunk_overlap = src.get("chunk_overlap")
        for p in sorted(root.rglob("*")):
            if not p.is_file():
                continue
            if p.name.lower() in CHATLOG_NAMES:
                ap = str(p.resolve())
                if ap in seen:
                    continue
                seen.add(ap)
                conversations = parse_chatlog(p)
                if conversations:
                    mtime = p.stat().st_mtime
                    seen_titles: set[str] = set()
                    for conv in conversations:
                        title, n = conv["title"], 2
                        while title.lower() in seen_titles:  # same-title convs would overwrite each other
                            title = f"{conv['title']} ({n})"
                            n += 1
                        seen_titles.add(title.lower())
                        conv["title"] = title
                        docs.append({"path": f"{ap}::{conv['title']}",
                                     "content": conv["text"],
                                     "hash": file_hash(conv["text"]),
                                     "tags": "chatlog", "links": "",
                                     "mtime": mtime,
                                     "chunk_size": src_chunk_size,
                                     "chunk_overlap": src_chunk_overlap})
                else:
                    print(f"[warn] unrecognized chat export, skipping: {p.name}")
                continue
            suffix = p.suffix.lower()
            if (suffix not in SUFFIXES and suffix not in (".pdf", ".docx")
                    and suffix not in HTML_SUFFIXES and suffix not in ORG_SUFFIXES):
                continue
            ap = str(p.resolve())
            if ap in seen:
                continue
            seen.add(ap)

            text = _read_text(p)
            if text is None or not text.strip():
                continue
            meta, body = parse_frontmatter(text)
            docs.append({"path": ap, "content": body, "hash": file_hash(text),
                         "tags": tags_of(meta), "links": extract_wikilinks(body),
                         "mtime": p.stat().st_mtime,
                         "chunk_size": src_chunk_size,
                         "chunk_overlap": src_chunk_overlap})
    return docs
